"""Score a run_baseline.py run against the intake-classification answer key.

Usage: python3 scripts/score.py out/runs/<model>.jsonl
Prints accuracy per field and writes out/runs/<model>-results.xlsx with three tabs:
Summary, All scenarios, and Not exactly right.

Answer key:
- Issue codes come from the "JSON case profiler" cell, which the dataset page calls the canonical
  gold label. Figures are also given for the rows where the "LIST Code" column agrees with it,
  since the review asked the Lab to confirm the others (the disputed rows).
- County is the Census code for the row's named jurisdiction, i.e. the corrected value from
  the review, not the dataset's own FIPS columns.
- Emergency comes from the "Emergency?" column.
"""
import json
import pathlib
import sys

from openpyxl import Workbook

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from build_review import table, text
from check_gold import CODE, STATES, collect, load_census, norm

ROOT = pathlib.Path(__file__).resolve().parent.parent
CSV = ROOT / "raw/intake-classification-inputs.csv"
TERMS = ROOT / "raw/list/terms.json"
CENSUS = ROOT / "raw/census/national_county2020.txt"

# Issue-code outcomes, best first. Each measure counts an outcome and every one above it.
EXACT, PARTIAL, LINE, BRANCH, AREA, WRONG = (
    "exact", "key's code, plus or minus others", "on the same path as the key's code",
    "right branch", "right area", "wrong area")
CODE_MEASURES = [
    ("Exact set of codes", EXACT, "The model's codes are exactly the key's codes."),
    ("At least one code right", PARTIAL, "One of the model's codes is one of the key's."),
    ("Same code or on the same path", LINE,
     "One of the model's codes is a key code, or sits above or below one on the same path in the taxonomy."),
    ("Right branch", BRANCH, "A model code shares its first two parts (e.g. HO-02) with a key code."),
    ("Right area", AREA, "A model code is in the same area of law (e.g. HO) as a key code."),
]
ORDER = [EXACT, PARTIAL, LINE, BRANCH, AREA, WRONG]


def stem(code):
    """The code without its trailing 00 levels: HO-02-00-00-00 -> HO-02."""
    parts = code.split("-")
    while len(parts) > 1 and parts[-1] == "00":
        parts.pop()
    return "-".join(parts)


def same_line(a, b):
    """True if one code is the other or sits above or below it in the taxonomy."""
    a, b = stem(a), stem(b)
    return a == b or a.startswith(b + "-") or b.startswith(a + "-")


def code_outcome(gold, got):
    if got == gold:
        return EXACT
    if not gold:
        # No key code means a non-legal referral: only an empty answer is right.
        return WRONG
    if got & gold:
        return PARTIAL
    if any(same_line(g, c) for g in gold for c in got):
        return LINE
    if {c[:5] for c in got} & {c[:5] for c in gold}:
        return BRANCH
    if {c[:2] for c in got} & {c[:2] for c in gold}:
        return AREA
    return WRONG


# Confounder types that the dataset page groups as "emergency-variation items".
EMERGENCY_VARIATIONS = {"mental_health_crisis", "ice_activity", "threats", "youth_at_risk"}
CATEGORIES = ("Ordinary", "Confounder", "Emergency variation", "False-positive emergency", "No legal issue")


def category(output):
    """The dataset's designed test categories, from the JSON's confounder_type."""
    kind = (output or {}).get("confounder_type") or ""
    if kind in EMERGENCY_VARIATIONS:
        return "Emergency variation"
    if kind.startswith("false_positive"):
        return "False-positive emergency"
    if kind.startswith("no_legal_issue"):
        return "No legal issue"
    return "Confounder" if kind else "Ordinary"


def judge(rows, preds, terms, disputed, by_name):
    """One dict per answered scenario: the key, the model's answer, and the outcome per field."""
    def named(codes):
        return "\n".join(f"{c}  {terms[c]['name'] if c in terms else '(not in LIST)'}" for c in sorted(codes))

    out = []
    for r in rows:
        sid = r["Scenario ID"]
        if sid not in preds:
            continue
        p = preds[sid]
        try:
            output = json.loads(r["JSON case profiler"])["output"]
            gold = set(CODE.findall(json.dumps(output.get("list_issue_codes") or [])))
        except json.JSONDecodeError:
            # One row's JSON does not parse; its column is the only copy of the key.
            output, gold = None, set(CODE.findall(r["LIST Code (from Legal Issue)"]))
        got = {c.strip().upper() for c in p.get("list_issue_codes") or []}
        state = STATES.get(r["Jdx it belongs to"].strip())
        p_state = (p.get("state") or "").strip().upper()
        truth = by_name.get((state, norm(r["Jurisdiction"]))) if state and r["Jurisdiction"].strip() else None
        key_emerg = r["Emergency?"].strip() == "checked"
        emerg = ("right" if p.get("is_emergency") == key_emerg
                 else "missed emergency" if key_emerg else "false alarm")
        out.append({
            "sid": sid, "channel": r["Input-Channel"], "query": r["User's Query"],
            "category": category(output), "kind": (output or {}).get("confounder_type") or "",
            "disputed": "yes" if sid in disputed else "",
            "key_codes": named(gold), "model_codes": named(got),
            "code": code_outcome(gold, got),
            "invented": sum(1 for c in got if c not in terms),
            "key_state": state or "", "model_state": p_state,
            "state": ("right" if p_state == state else "wrong") if state else "not scored (no key)",
            "key_county": r["Jurisdiction"].strip(), "model_county": (p.get("county_name") or "").strip(),
            "county": ("right" if by_name.get((p_state, norm(p.get("county_name") or ""))) == truth else "wrong")
                      if truth else "not scored (no key)",
            "key_fips": truth or "", "model_fips": (p.get("county_fips") or "").strip(),
            "fips": ("right" if (p.get("county_fips") or "").strip() == truth else "wrong")
                    if truth else "not scored (no key)",
            "key_emerg": "yes" if key_emerg else "no", "model_emerg": "yes" if p.get("is_emergency") else "no",
            "emerg": emerg,
        })
    return out


def summarise(res):
    """Rows of (field, measure, right, out of, percent, meaning)."""
    def line(field, measure, right, total, meaning):
        return (field, measure, right, total, f"{100 * right / total:.0f}%" if total else "", meaning)

    out = []
    for field, coded in (("Issue code", res),
                         ("Issue code, undisputed rows only", [x for x in res if not x["disputed"]])):
        out += [line(field, name, sum(1 for x in coded if ORDER.index(x["code"]) <= ORDER.index(worst)),
                     len(coded), meaning) for name, worst, meaning in CODE_MEASURES]
    for field, key, meaning in (
            ("State", "state", "The model's state is the key's state."),
            ("County, by name", "county", "The county the model named is the key's county."),
            ("County, by FIPS code", "fips", "The five-digit code the model gave is the Census code."),
            ("Emergency", "emerg", "The model's yes/no matches the key's.")):
        scored = [x for x in res if x[key] in ("right", "wrong", "missed emergency", "false alarm")]
        out.append(line(field, "Correct", sum(1 for x in scored if x[key] == "right"), len(scored), meaning))
    # One accuracy figure hides which way the emergency mistakes run, so give both directions.
    caught = sum(1 for x in res if x["key_emerg"] == "yes" and x["model_emerg"] == "yes")
    out.append(line("Emergency", "Key's emergencies the model flagged (recall)", caught,
                    sum(1 for x in res if x["key_emerg"] == "yes"), "Low means emergencies are being missed."))
    out.append(line("Emergency", "Model's flags that the key agrees with (precision)", caught,
                    sum(1 for x in res if x["model_emerg"] == "yes"), "Low means false alarms."))
    return out


def by_category(res):
    """Rows of (category, scenarios, then right-counts for the measures that matter most)."""
    out = []
    for cat in CATEGORIES:
        xs = [x for x in res if x["category"] == cat]
        out.append((cat, len(xs),
                    sum(1 for x in xs if ORDER.index(x["code"]) <= ORDER.index(PARTIAL)),
                    sum(1 for x in xs if ORDER.index(x["code"]) <= ORDER.index(LINE)),
                    sum(1 for x in xs if x["county"] == "right"),
                    sum(1 for x in xs if x["emerg"] == "right"),
                    sum(1 for x in xs if x["emerg"] == "missed emergency"),
                    sum(1 for x in xs if x["emerg"] == "false alarm")))
    return out


CATEGORY_HEAD = ["Category", "Scenarios", "At least one code right", "Same code or on the same path",
                 "County right", "Emergency right", "Emergencies missed", "False alarms"]


def workbook(path, run_path, res, summary, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    text(ws, 1, f"Intake classification baseline: {run_path.stem}", bold=True, size=14)
    details = [
        ("Model", meta.get("model", run_path.stem)),
        ("Run date", meta.get("date", "not recorded")),
        ("Dataset file", f'{CSV.name} (sha256 {meta.get("dataset_sha256", "not recorded")[:12]})'),
        ("Scenarios answered", f"{len(res)} of {meta.get('scenarios', 'unknown')}"),
        ("Rows where the LIST Code column disagrees with the JSON key", sum(1 for x in res if x["disputed"])),
        ("Codes the model gave that are not in LIST", sum(x["invented"] for x in res)),
        ("Emergencies missed / false alarms", f'{sum(1 for x in res if x["emerg"] == "missed emergency")} / '
                                               f'{sum(1 for x in res if x["emerg"] == "false alarm")}'),
    ]
    for i, (k, v) in enumerate(details, 3):
        text(ws, i, k, bold=True)
        ws.cell(row=i, column=2, value=str(v))
    top = 4 + len(details)
    table(ws, top, ["Field", "Measure", "Right", "Out of", "Percent", "What counts as right"],
          summary, [34, 28, 8, 8, 9, 80])
    row = top + len(summary) + 2
    cats = by_category(res)
    table(ws, row, CATEGORY_HEAD, cats, [34, 10, 12, 14, 10, 12, 12, 10])
    ws.freeze_panes = None
    row += len(cats) + 2
    text(ws, row, "Instructions given to the model (followed by the LIST code list)", bold=True)
    text(ws, row + 1, meta.get("instructions", "not recorded"), width_cols=6)
    ws.row_dimensions[row + 1].height = 260
    text(ws, row + 3, "Message format: " + meta.get("message_format", "not recorded"))

    cols = [("Scenario", "sid", 16), ("Category", "category", 16), ("Confounder type", "kind", 22),
            ("Key disputed", "disputed", 9), ("Channel", "channel", 11), ("Message", "query", 70),
            ("Key: issue codes", "key_codes", 44), ("Model: issue codes", "model_codes", 44),
            ("Issue code result", "code", 24),
            ("Key: state", "key_state", 8), ("Model: state", "model_state", 8), ("State", "state", 10),
            ("Key: county", "key_county", 20), ("Model: county", "model_county", 20), ("County", "county", 10),
            ("Key: FIPS", "key_fips", 9), ("Model: FIPS", "model_fips", 9), ("FIPS", "fips", 10),
            ("Key: emergency", "key_emerg", 10), ("Model: emergency", "model_emerg", 10),
            ("Emergency", "emerg", 17)]
    ws = wb.create_sheet("All scenarios")
    table(ws, 1, [c[0] for c in cols], [[x[c[1]] for c in cols] for x in res], [c[2] for c in cols])

    wrong = []
    for x in res:
        for field, key, want, got in (("issue code", "code", "key_codes", "model_codes"),
                                      ("state", "state", "key_state", "model_state"),
                                      ("county", "county", "key_county", "model_county"),
                                      ("county FIPS", "fips", "key_fips", "model_fips"),
                                      ("emergency", "emerg", "key_emerg", "model_emerg")):
            if x[key] not in ("right", EXACT) and not x[key].startswith("not scored"):
                wrong.append([x["sid"], field, "" if x[key] == "wrong" else x[key], x[want], x[got], x["query"]])
    ws = wb.create_sheet("Not exactly right")
    table(ws, 1, ["Scenario", "Field", "Kind", "Key", "Model said", "Message"], wrong, [16, 12, 30, 44, 44, 70])
    wb.save(path)
    return len(wrong)


def main(run_path):
    run_path = pathlib.Path(run_path)
    preds = {}
    for line in open(run_path):
        rec = json.loads(line)
        if rec["pred"] is not None:
            preds[rec["sid"]] = rec["pred"]
    meta_path = run_path.with_suffix(".meta.json")
    meta = json.load(open(meta_path)) if meta_path.exists() else {}
    rows, terms, _, disputes, _ = collect(CSV, TERMS, CENSUS)
    meta["scenarios"] = len(rows)
    by_name, _ = load_census(CENSUS)
    res = judge(rows, preds, terms, {d["sid"] for d in disputes}, by_name)
    summary = summarise(res)

    print(f"{run_path.stem}: {len(res)} scenarios scored\n")
    if len(res) < len(rows):
        print(f"  WARNING: {len(rows) - len(res)} scenarios have no answer and are not in any figure below.\n")
    for field, measure, right, total, percent, _ in summary:
        print(f"  {field + ': ' + measure:<42} {right}/{total} ({percent})")
    print()
    print("  " + " | ".join(CATEGORY_HEAD))
    for c in by_category(res):
        print("  " + " | ".join(str(v) for v in c))
    out = run_path.with_name(run_path.stem + "-results.xlsx")
    n = workbook(out, run_path, res, summary, meta)
    print(f"\n{n} answers not exactly right. Workbook written to {out}")


if __name__ == "__main__":
    main(sys.argv[1])
